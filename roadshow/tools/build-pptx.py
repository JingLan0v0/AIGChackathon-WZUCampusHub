#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
路演 PPT 生成器（14 页 · 页面内点击逐条展现）
==========================================================
为什么不用现成的「一键生成 PPT」：
  它们的产物是「静态整页」，做不到你要的「一页里点一下出一块」。
  这里用 python-pptx 精确控制每一个形状，再把「点击触发出场动画」
  以 OOXML（<p:timing>）手工注入到每页 —— 这是 PowerPoint 原生动画格式，
  效果和你在 PowerPoint 里手动加「出现（单击时）」完全一样。

用法：
  python roadshow/tools/build-pptx.py
产物：
  roadshow/校园百事通-路演.pptx
"""

import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, 'shots')
OUT = os.path.join(ROOT, '校园百事通-路演.pptx')

FONT = '微软雅黑'

# 配色：跟着「校园百事通」站点走，演示时视觉一致
PRIMARY = '3A5BFF'
PRIMARY_D = '2A44D6'
PRIMARY_L = 'E8EDFF'
VIOLET = '7A5CFF'
VIOLET_L = 'F0EBFF'
TEXT = '131A2A'
TEXT2 = '45506A'
MUTED = '8892A6'
LINE = 'E1E8F3'
SURF = 'FFFFFF'
SURF2 = 'F4F7FC'
RED = 'F04438'
RED_L = 'FDECEA'
GREEN = '12B76A'
GREEN_L = 'E6F8EF'
YELLOW = 'F79009'
YELLOW_L = 'FEF4E6'
DARK = '16204A'
DARK2 = '23307A'

EMU_IN = 914400
SLIDE_W = 13.333
SLIDE_H = 7.5

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]

# 每页的点击分组：slide -> [[shape, shape], [shape], ...]
GROUPS = {}


# ---------------------------------------------------------------- 基础工具
def new_slide(bg=SURF, groups=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    GROUPS[len(prs.slides) - 1] = groups if groups is not None else []
    return s


def rect(slide, x, y, w, h, fill=None, line=None, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(
        shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
        sh.line.width = Pt(0.9)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.word_wrap = True
    return sh


def runs(*items, **kw):
    """一行里的多个片段：runs(("文字", 16, True, TEXT), ("→ 补充", 14, False, RED))"""
    return {'runs': list(items), 'align': kw.get('align', PP_ALIGN.LEFT),
            'line': kw.get('line', 1.15), 'before': kw.get('before', 0)}


def write(shape, paras, anchor=MSO_ANCHOR.MIDDLE, mt=0.06, ml=0.14):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_top = tf.margin_bottom = Inches(mt)
    tf.margin_left = tf.margin_right = Inches(ml)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        para.line_spacing = p.get('line', 1.15)
        if p.get('before'):
            para.space_before = Pt(p['before'])
        for r in p['runs']:
            text, size, bold, color = r[0], r[1], r[2], r[3]
            run = para.add_run()
            run.text = text
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)
            run.font.name = FONT
    return shape


def textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    return write(tb, paras, anchor=anchor, mt=0.02, ml=0.02)


def card(slide, x, y, w, h, paras, fill=SURF2, line=LINE, radius=0.10,
         anchor=MSO_ANCHOR.MIDDLE):
    sh = rect(slide, x, y, w, h, fill=fill, line=line, radius=radius)
    return write(sh, paras, anchor=anchor)


def title_bar(slide, kicker, title, accent=PRIMARY):
    """页眉：彩色小标签 + 大标题 + 顶部细色条"""
    rect(slide, 0, 0, SLIDE_W, 0.075, fill=accent, shape=MSO_SHAPE.RECTANGLE)
    pill = rect(slide, 0.62, 0.42, 3.6, 0.34, fill=PRIMARY_L, radius=0.5)
    write(pill, [runs((kicker, 11.5, True, PRIMARY), align=PP_ALIGN.CENTER)], mt=0.02, ml=0.06)
    textbox(slide, 0.58, 0.88, 12.2, 0.72,
            [runs((title, 27, True, TEXT))], anchor=MSO_ANCHOR.TOP)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def phone(slide, name, x, y, h, crop_bottom=0.0):
    p = os.path.join(SHOTS, name + '.png')
    pic = slide.shapes.add_picture(p, Inches(x), Inches(y), height=Inches(h))
    if crop_bottom:
        pic.crop_bottom = crop_bottom
    return pic


def page_no(slide, n, color=MUTED):
    textbox(slide, 12.4, 7.02, 0.6, 0.3, [runs((str(n), 11, False, color), align=PP_ALIGN.RIGHT)],
            anchor=MSO_ANCHOR.MIDDLE)


# ================================================================ 01 封面
s = new_slide(bg=PRIMARY)
rect(s, 8.4, -1.2, 6.6, 6.6, fill='4C6BFF', radius=0.35)
rect(s, -1.6, 4.4, 7.2, 5.0, fill='3352E8', radius=0.35)
logo = rect(s, 0.9, 0.85, 0.92, 0.92, fill='5C79FF', radius=0.28)
write(logo, [runs(('🎓', 26, False, 'FFFFFF'), align=PP_ALIGN.CENTER)], mt=0.02, ml=0.02)
textbox(s, 1.98, 0.95, 8.4, 0.5,
        [runs(('温州大学 · 校园生活百事通 · 赛道二', 14, True, 'FFFFFF'))])
textbox(s, 0.86, 2.05, 11.6, 1.5, [runs(('校园百事通', 58, True, 'FFFFFF'))])
sub = textbox(s, 0.9, 3.35, 11.0, 0.6,
              [runs(('温大学生的一站式校园助手', 22, True, 'EAF0FF'))])
line1 = textbox(s, 0.9, 4.05, 11.4, 1.1, [
    runs(('把「天气 · 课表 · 调课通知 · 作业 DDL · 课外计划 · 二手 · 失物 · 问答」', 16, False, 'D8E1FF')),
    runs(('收进一个入口、一次登录，而且听得懂人话。', 16, False, 'D8E1FF'), before=4),
])
foot = textbox(s, 0.9, 5.75, 11.4, 1.1, [
    runs(('选手：forey　|　温州大学计算机与人工智能学院', 13, False, 'D8E1FF')),
    runs(('现场演示：campus-hub-57981.app.workbuddy.host　（演示账号 20230427）', 13, False, 'FFFFFF'), before=3),
    runs(('▶ 点击鼠标，本页内容会一条条展开', 12, True, 'FFFFFF'), before=6),
])
GROUPS[0] = [[sub], [line1], [foot]]
notes(s, '先自报家门：温州大学 · 校园百事通 · 赛道二。一句话说明白我们做的是什么——把学生每天要开的那几个工具，收进一个入口。')

# ================================================================ 02 痛点
s = new_slide()
title_bar(s, '01 · 问题', '查一次校园信息，要付出多少步？')
PAINS = [
    ('☀️', '想知道今天穿什么', '得再开一个天气 App / 网页'),
    ('📝', '想记住作业 DDL', '打开备忘录，自己再敲一遍'),
    ('📅', '想看课表', '先过统一身份认证，再在一堆菜单里找入口'),
    ('🔔', '想确认有没有调课', '翻教务系统，翻不到就当没发生'),
    ('🛒', '想买二手 / 找失物', '微信群爬楼，刷过去就没了'),
]
y = 1.78
pains = []
for ico, main, then in PAINS:
    c = card(s, 0.62, y, 12.1, 0.78, [
        runs((ico + '  ', 15, False, TEXT), (main, 16, True, TEXT),
             ('　→　' + then, 14.5, False, RED)),
    ], anchor=MSO_ANCHOR.MIDDLE)
    pains.append(c)
    y += 0.9
bar = card(s, 0.62, 6.32, 12.1, 0.78, [
    runs(('五件事 = 五个入口 = 五次登录与切换', 19, True, RED), align=PP_ALIGN.CENTER)],
    fill=RED_L, line='F7C9C4')
GROUPS[1] = [[c] for c in pains] + [[bar]]
notes(s, '不要急着讲产品，先讲他每天的真实处境。每点一下出一个场景，语气可以稍带吐槽——评委自己也是这么过来的。最后一句「五件事 = 五个入口」是转折点，说完停一秒再翻页。')
page_no(s, 2)

# ================================================================ 03 归纳
s = new_slide()
title_bar(s, '02 · 问题', '问题不在「功能不够」，在「入口太散」')
ITEMS3 = [
    ('1', '入口散', '一天里在 5 个工具之间来回切，每次都重新找、重新登', PRIMARY, SURF2, LINE),
    ('2', '会漏', '调课通知躺在教务系统里，不主动看就错过——去了才发现教室换了', YELLOW, YELLOW_L, 'F8DFB8'),
    ('3', '记不住', '作业靠脑子记，备忘录又懒得敲一遍——工具越多，越不想用', VIOLET, SURF2, LINE),
]
y = 1.85
g3 = []
for num, name, desc, col, bg, ln in ITEMS3:
    c = card(s, 0.62, y, 12.1, 1.12, [], fill=bg, line=ln)
    badge = rect(s, 0.86, y + 0.31, 0.5, 0.5, fill=col, radius=0.28)
    write(badge, [runs((num, 15, True, 'FFFFFF'), align=PP_ALIGN.CENTER)], mt=0.02, ml=0.02)
    write(c, [runs((name, 19, True, TEXT)), runs((desc, 14, False, TEXT2), before=3)], ml=0.85)
    g3.append(c)
    y += 1.26
concl = textbox(s, 0.62, 5.85, 12.1, 0.8,
                [runs(('学生要的不是更多功能，是一个「不用记、不用切」的地方。', 21, True, PRIMARY))])
GROUPS[2] = [[c] for c in g3] + [[concl]]
notes(s, '把上一页的具体场景收成三个抽象问题：散、漏、记不住。这三条后面正是我们的三个卖点（统一入口 / 主动提醒 / 说人话），这里要点透，后面才能呼应。')
page_no(s, 3)

# ================================================================ 04 答案
s = new_slide()
title_bar(s, '03 · 我们的答案', '一个入口，一次登录，说人话就行')
ANS = [
    ('🧩', '九个入口，一个顶栏', '首页 / 规划 / 课表 / 调课 / 作业 / 二手 / 失物 / 问答 / 我的，全部一个顶栏切到底', SURF2, LINE),
    ('🔑', '一次登录，跨页免登', '学号绑定后本机所有页面直进；多账号可一键切换，数据按学号互相隔离', SURF2, LINE),
    ('✨', '说人话就行', '「明晚 7 点图书馆自习」「明天上午请病假，因为发烧」——一句话，它自己判断该记成什么', PRIMARY_L, 'CBD7FF'),
]
g4 = []
for i, (ico, name, desc, bg, ln) in enumerate(ANS):
    x = 0.62 + i * 4.13
    c = card(s, x, 2.0, 3.9, 2.75, [
        runs((ico, 26, False, TEXT)),
        runs((name, 18, True, TEXT), before=6),
        runs((desc, 13, False, TEXT2), before=5, line=1.3),
    ], fill=bg, line=ln, anchor=MSO_ANCHOR.TOP)
    g4.append(c)
closing = textbox(s, 0.62, 5.25, 12.1, 0.9, [
    runs(('我们不是把功能搬到一起，而是让它', 19, False, TEXT2),
         ('听得懂人话、自己找到该去的位置', 19, True, PRIMARY), ('。', 19, False, TEXT2))])
GROUPS[3] = [[c] for c in g4] + [[closing]]
notes(s, '转折页。我们的答案就是三个词：统一入口、一次登录、说人话。第三个是关键差异——把功能堆在一起的产品很多，但让学生用一句人话说事的不多，这也是 AI 真正上场的地方。')
page_no(s, 4)

# ================================================================ 05 九个模块
s = new_slide()
title_bar(s, '04 · 它长这样', '九个模块，覆盖校园生活的一天')
MODS = [
    ('🏠', '首页', '天气 + 出行建议 + 今日规划摘要'),
    ('🧭', '今日规划', '按空档与 DDL 排出的专注安排'),
    ('📅', '智能课表', '温大真实作息 11 节 · 支持导入'),
    ('🔔', '调课通知', '抓取 → 解析 → 一键应用到课表'),
    ('📝', '作业 DDL', '倒计时 / 逾期提醒 / 一键完成'),
    ('🛒', '二手集市', '发布 · 筛选 · 联系卖家'),
    ('🔍', '失物招领', '丢了 / 捡到，分开找'),
    ('💬', '校园问答', '选课、食堂、报修都有人答'),
    ('👤', '我的', '学籍 · 导入导出 · 多账号切换'),
]
cw, ch, gap = 2.38, 1.32, 0.16
cells = []
for i, (ico, name, desc) in enumerate(MODS):
    cx = 0.62 + (i % 3) * (cw + gap)
    cy = 1.82 + (i // 3) * (ch + gap)
    c = card(s, cx, cy, cw, ch, [
        runs((ico + ' ', 14, False, TEXT), (name, 14.5, True, TEXT)),
        runs((desc, 10.5, False, MUTED), before=3, line=1.2),
    ], anchor=MSO_ANCHOR.TOP)
    cells.append(c)
ph = phone(s, '02-home', 10.35, 1.82, 4.4, crop_bottom=0.06)
GROUPS[4] = [[cells[0], cells[1], cells[2]], [cells[3], cells[4], cells[5]],
             [cells[6], cells[7], cells[8]], [ph]]
notes(s, '快速过一遍产品形态，不用逐个讲。左边九宫格分三批亮，右边是真机首页截图——让评委先看到一个真实产品，而不是架构图。')
page_no(s, 5)

# ================================================================ 06-09 四张功能页
def feature_slide(kicker, title, shot, crop, items, tail, tail_style,
                  note_text, n, ih=0.86):
    """ih = 单条卡片高度：文案长的页（比如课表那页）调高一点，避免文字顶出框"""
    s = new_slide()
    title_bar(s, kicker, title)
    phone(s, shot, 0.72, 1.72, 4.55, crop_bottom=crop)
    gs = []
    y = 1.72
    for tag, t, d in items:
        c = card(s, 3.42, y, 9.3, ih, [])
        pill = rect(s, 3.58, y + 0.19, 0.62, 0.3, fill=VIOLET_L, radius=0.45)
        write(pill, [runs((tag, 9.5, True, VIOLET), align=PP_ALIGN.CENTER)], mt=0.01, ml=0.02)
        write(c, [runs((t, 15, True, TEXT)),
                  runs((d, 12.5, False, TEXT2), before=2, line=1.2)], ml=0.76)
        gs.append(c)
        y += ih + 0.10
    fill, ln = tail_style
    tail_sh = card(s, 3.42, y + 0.04, 9.3, 0.86, [
        runs((tail, 12.5, False, TEXT2), line=1.25)], fill=fill, line=ln)
    GROUPS[n] = [[c] for c in gs] + [[tail_sh]]
    notes(s, note_text)
    page_no(s, n + 1)
    return s

feature_slide(
    '05 · 核心操作一', '一句话记下来：说人话，它自己归类',
    '03-nl-parse', 0.02,
    [('拆句', '一句话拆成两件事', '分号 / 换行都能拆，最多一次记 5 件，互不干扰'),
     ('分类', '自动判断该记成什么', '作业 DDL / 课外计划 / 课表补充 / 请假条，四类四种去处'),
     ('抽字段', '时间、地点、类型、重要度全抽出来', '「高数」自动对上真实课表里的「高等数学 A(一)」'),
     ('可改', '先给你看「我理解为」，确认才写库', '字段全都能现场改，不满意可以「不要这条」')],
    '看不懂就直说看不懂：认不出来的字段会明确标出让用户补，不会瞎填一个默认值糊过去。',
    (YELLOW_L, 'F8DFB8'),
    '这是核心交互，讲慢一点。强调三点：① 一句话里可以有多件事，它自己拆；② 理解结果先给用户核对、可改，不是黑盒直接写库；③ 认不出来会直说没看懂，绝不瞎建条目。',
    5)

feature_slide(
    '06 · 核心操作二', '请假条：一句话，生成能直接发的正文',
    '04-leave', 0.30,
    [('识别', '假别 / 时间 / 时段 / 天数', '病假 · 明天上午（半天），也认「请事假两天」「周五第 3-4 节」'),
     ('取理由', '理由从「因为」后面取，多余的字掐掉', '说「因为发烧要去校医院」，正文里就是这一句'),
     ('成文', '自动拼出正式请假条', '抬头 · 学院班级学号 · 假别时间 · 补课承诺 · 落款日期')],
    '不写进日程，只给你成品：请假条是发给人的，不是日程；内容只在本机生成，姓名学院没填就留占位，系统不替你编。',
    (GREEN_L, 'BFE9D3'),
    '最能体现「AI 省事」的一页：请假条本来是学生最烦的格式活，现在一句话生成、一键复制。特别提一句「不写进日程、不上传、缺信息留占位」——这是我们对隐私和诚实的态度。',
    6)

feature_slide(
    '07 · 核心操作三', '智能课表：按温大真实作息长出来的',
    '05-schedule', 0.02,
    [('真实', '11 节真实作息，不是拍脑袋的模板', '上午 1–5 节（第 5 节 11:45–12:25 午前最后一节，不排课）· 下午从第 6 节 13:30 起 · 晚上为第 11 节'),
     ('导入', '三种方式把真实课表弄进来', '课表 PDF 识别 / 同源抓取脚本一键导入 / 一句话补充'),
     ('细节', '连堂、周次、教室、教师都算对', '「第 2-9 节」不会被截断；同课同色；支持「南5-A101」这类门牌')],
    '空课表也不是死路：新同学第一次进来，页面直接给「去统一身份认证查课表 → 回来粘贴导入」的三步引导。',
    (PRIMARY_L, 'CBD7FF'),
    '课表是地基。我们把温大真实作息做成单一数据源：上午 1-5 节、下午从第 6 节 13:30 起、晚上一整段算第 11 节。讲师门难点时可以说「第 5 节 11:45 是午前最后一节，学校不排课」——这种细节最能证明我们真的对着温大做的。',
    7, ih=1.06)

feature_slide(
    '08 · 核心操作四', '调课通知：抓回来，自动落到课表上',
    '06-notice', 0.02,
    [('抓取', '从教务系统把通知原文取回来', '自己写的同源脚本一键复制 → 回页面粘贴导入；不碰密码、不调校外接口'),
     ('解析', '把一段通知读成「哪门课、怎么变」', '原时间 → 新时间、新教室、停课/补课，识别「第 3 周周二第 6-7 节」'),
     ('匹配', '自动对上你课表里的那门课', '按课程名 + 星期 + 起始节匹配，首页提示「这周先别按老时间去」'),
     ('应用', '一键应用到课表，原安排标灰保留', '可撤销、可标记已读，影响今天的通知提到最上面')],
    '不编造：没有导入过真实通知的学号，页面一条假通知都不给，只给导入引导（只有演示账号带样本）。',
    (YELLOW_L, 'F8DFB8'),
    '这条解决「会漏」：通知从教务系统抓回来，解析成结构化，再自动匹配课表里的课。注意我们很克制——真实学号看不到任何编造样本，没有就是没有。这种诚实反而是加分项。',
    8)

# ================================================================ 10 AI 在做什么
s = new_slide()
title_bar(s, '09 · AI 到底做了什么', 'AI 不是在贴标签，它在替学生做这几件事')
LEFT = [
    ('理解', '自然语言 → 结构化数据', '从「明晚 7 点图书馆自习 2 小时」里抽出日期、时刻、时长、地点、类型、重要度'),
    ('分类', '意图识别', '判断这是计划、作业 DDL、课表补充，还是请假条——四类四种处理路径'),
    ('生成', '请假条正文生成', '按学生信息、课时、理由、天数组织成正式文书，可直接发送'),
    ('解析', '教务通知理解与课表改写', '从自然语言通知里定位课程、时间、教室变化，映射到课表具体节次'),
]
RIGHT = [
    ('识别', '课表 PDF 识别', '读教务系统导出的课表 PDF，把「8:20–9:00」映射到第 1 节，连堂自动合并'),
    ('建议', '天气 → 出行建议', '带不带伞、穿什么、骑车还是坐校车、几点适合运动'),
    ('排程', '今日规划自动排', '在课表空档里按 DDL 与重要度排出「现在最该做什么」，并解释为什么'),
]
g10 = []
for col, items in ((0, LEFT), (1, RIGHT)):
    y = 1.78
    for tag, t, d in items:
        x = 0.62 + col * 6.31
        c = card(s, x, y, 6.1, 0.98, [])
        pill = rect(s, x + 0.16, y + 0.2, 0.6, 0.28, fill=VIOLET_L, radius=0.45)
        write(pill, [runs((tag, 9.5, True, VIOLET), align=PP_ALIGN.CENTER)], mt=0.01, ml=0.02)
        write(c, [runs((t, 14, True, TEXT)),
                  runs((d, 11.5, False, TEXT2), before=2, line=1.18)], ml=0.72)
        g10.append(c)
        y += 1.1
note10 = card(s, 0.62, 6.28, 12.1, 0.72, [
    runs(('共同点：', 12.5, True, TEXT2),
         ('都在替代学生「自己想、自己敲、自己翻」的动作，而不是多给一个界面。', 12.5, False, TEXT2))],
    fill=PRIMARY_L, line='CBD7FF')
GROUPS[9] = [[c] for c in g10] + [[note10]]
notes(s, '评委最关注「AI 含量」的一页。一条条点亮，每条配一句「它替代了什么人工动作」。时间紧就重点讲前三条：自然语言理解 / 请假条生成 / 通知解析。')
page_no(s, 10)

# ================================================================ 11 为什么本地引擎
s = new_slide()
title_bar(s, '10 · 技术选择', '为什么是本地智能引擎，而不是调一个大模型接口？')
WHY = [
    ('①', '断网也能用', '现场演示不赌网络。整个应用是零依赖静态站，双击本地文件就能跑。', SURF2, LINE),
    ('②', '隐私留在本机', '课表、作业、请假条都是私事，不该往第三方接口发。数据全部存在本机。', SURF2, LINE),
    ('③', '每一步都能讲清楚', '为什么把它读成 19 点？规则摆在明面上，界面上也能解释，不是黑盒。', SURF2, LINE),
    ('④', '可替换，不是封闭', '解析能力收在 NL.parseAll() / Adjust.sync() 这样的单一入口，接大模型只需换一个函数。', PRIMARY_L, 'CBD7FF'),
]
g11 = []
for i, (num, t, d, bg, ln) in enumerate(WHY):
    x = 0.62 + (i % 2) * 6.31
    y = 1.95 + (i // 2) * 1.62
    c = card(s, x, y, 6.1, 1.44, [
        runs((num + '  ', 17, True, PRIMARY), (t, 17, True, TEXT)),
        runs((d, 12.5, False, TEXT2), before=5, line=1.25),
    ], fill=bg, line=ln, anchor=MSO_ANCHOR.TOP)
    g11.append(c)
tail11 = textbox(s, 0.62, 5.5, 12.1, 0.7,
                 [runs(('我们先用规则把高频表达做扎实，再谈要不要上大模型——顺序不能反。', 16, False, TEXT2))])
GROUPS[10] = [[c] for c in g11] + [[tail11]]
notes(s, '主动交代技术选择，评委会问「为什么不用大模型」。答：不是不会用，是这四个理由。最后一条要强调——架构是可替换的，接上大模型能力只会更强。')
page_no(s, 11)

# ================================================================ 12 教务接入
s = new_slide()
title_bar(s, '11 · 真实数据怎么进来', '不硬闯登录，用学生自己的会话把数据带回来')
warn = card(s, 0.62, 1.8, 12.1, 0.95, [
    runs(('先说难点：', 14.5, True, TEXT),
         ('统一身份认证是网易 SSO，表单动态渲染、密码用公钥加密、还有人机验证——静态页面直连拿不到会话，硬做就是死路。', 14, False, TEXT2),
         )], fill=YELLOW_L, line='F8DFB8')
STEPS = [
    ('① 同源抓取', '学生在自己已登录的教务页面上运行我们给的脚本（一行复制），把课表与通知原文读出来'),
    ('② 一键导入', '回应用粘贴，本地解析成课表与通知；也可以直接上传课表 PDF'),
    ('③ 同一条链路', '导入数据与演示样本走完全相同的解析逻辑，正式版只需换掉取数那一步'),
]
g12 = []
for i, (t, d) in enumerate(STEPS):
    x = 0.62 + i * 4.13
    c = card(s, x, 2.95, 3.9, 1.72, [
        runs((t, 16, True, TEXT)),
        runs((d, 12.5, False, TEXT2), before=6, line=1.3),
    ], fill=PRIMARY_L, line='CBD7FF', anchor=MSO_ANCHOR.TOP)
    g12.append(c)
tail12 = textbox(s, 0.62, 4.95, 12.1, 0.7, [
    runs(('全程', 17, False, TEXT2), ('不接触学生密码', 17, True, GREEN),
         ('、不调校外接口、不落地到任何服务器。', 17, False, TEXT2))])
GROUPS[11] = [[warn]] + [[c] for c in g12] + [[tail12]]
notes(s, '这页讲「我们没吹牛」：统一身份认证是网易 SSO + 人机验证，静态页直连拿不到会话，所以我们不硬闯，改成同源抓取脚本 + 一键导入。把自己的技术边界说清楚，比嘴上说「全面打通」更可信。')
page_no(s, 12)

# ================================================================ 13 演示动线
s = new_slide()
title_bar(s, '12 · 现场演示路线', '跟着点一遍，30 秒看清全貌')
STEPS13 = [
    ('1', '打开链接 → 学号登录', '演示账号 20230427 · 首页直接看到天气、出行建议与今日规划摘要'),
    ('2', '一句话记下来', '说「高数作业 周五 23:59 交；明晚 7 点图书馆自习 2 小时」→ 出可改的解析卡 → 确认落库'),
    ('3', '请假条生成', '说「明天上午请病假，因为发烧要去校医院」→ 生成正文 → 复制（可现场粘贴到聊天框）'),
    ('4', '课表 + 调课通知', '看温大真实作息 11 节；调课页点「应用到课表」，回课表页看它变了'),
    ('5', '手机打开同一个链接', '换一个学号登录 → 数据互相隔离；二手 / 失物 / 问答也能直接用'),
]
g13 = []
y = 1.82
for num, t, d in STEPS13:
    c = card(s, 0.62, y, 12.1, 0.92, [])
    badge = rect(s, 0.86, y + 0.24, 0.44, 0.44, fill=PRIMARY, radius=0.3)
    write(badge, [runs((num, 13, True, 'FFFFFF'), align=PP_ALIGN.CENTER)], mt=0.01, ml=0.01)
    write(c, [runs((t, 15.5, True, TEXT)),
              runs(('　' + d, 12.5, False, TEXT2))], ml=0.8)
    g13.append(c)
    y += 1.02
GROUPS[12] = [[c] for c in g13]
notes(s, '给评委一个「接下来看什么」的路线图，方便你现场操作时他们跟得上。这一页也可以当彩排清单——照着点一遍，30 秒内讲完。')
page_no(s, 13)

# ================================================================ 14 收尾
s = new_slide(bg=DARK)
rect(s, 8.6, -1.4, 6.4, 6.4, fill=DARK2, radius=0.35)
pill = rect(s, 0.9, 0.8, 3.2, 0.34, fill=DARK2, radius=0.5)
write(pill, [runs(('一句话收尾', 11.5, True, 'FFFFFF'), align=PP_ALIGN.CENTER)], mt=0.02, ml=0.06)
big = textbox(s, 0.86, 1.5, 11.6, 1.9, [
    runs(('把五个入口，', 34, True, 'FFFFFF')),
    runs(('收进一个听得懂人话的入口', 34, True, '9DB2FF'), before=6)])
desc = textbox(s, 0.9, 3.45, 11.4, 1.0, [
    runs(('天气、课表、调课、作业、二手失物——过去要开五个工具、登五次录、记五遍；', 15, False, 'C7D2FF')),
    runs(('现在一句话说给它听，它自己判断该去哪、该记成什么、该提醒什么。', 15, False, 'C7D2FF'), before=3)])
done = card(s, 0.62, 4.72, 6.1, 1.5, [
    runs(('已完成', 15, True, 'FFFFFF')),
    runs(('9 个模块全通 · 温大真实作息 · 多账号隔离 · 课表 PDF 识别与导入 · 自然语言记录 · 请假条生成', 12, False, 'D8E1FF'), before=6, line=1.3),
], fill=DARK2, line='3E4E9E', anchor=MSO_ANCHOR.TOP)
nextc = card(s, 6.62, 4.72, 6.1, 1.5, [
    runs(('下一步', 15, True, 'FFFFFF')),
    runs(('接真实教务接口（课表 / 成绩 / 通知）· 二手交易闭环 · 校园问答接入学校知识库', 12, False, 'D8E1FF'), before=6, line=1.3),
], fill=DARK2, line='3E4E9E', anchor=MSO_ANCHOR.TOP)
foot14 = textbox(s, 0.9, 6.45, 11.4, 0.6, [
    runs(('🎓 校园百事通　', 15, True, 'FFFFFF'),
         ('campus-hub-57981.app.workbuddy.host', 14, False, '9DB2FF'))])
GROUPS[13] = [[desc], [done, nextc], [foot14]]
notes(s, '收尾不要念清单，只讲一句核心话术：「把五个入口，收进一个听得懂人话的入口。」然后留链接请评委自己试。')

# ================================================================ 注入点击动画
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def timing_xml(groups):
    """生成「单击依次出现」的 <p:timing>（PowerPoint 原生格式）"""
    ctn = [10]  # 自增 id

    def nid():
        ctn[0] += 1
        return ctn[0]

    def eff(shape, node_type):
        return (
            '<p:par><p:cTn id="%d" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" '
            'grpId="0" nodeType="%s"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            '<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold">'
            '<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
            '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
            '</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            '</p:childTnLst></p:cTn></p:par>'
        ) % (nid(), node_type, nid(), shape.shape_id)

    par_list = []
    for gi, group in enumerate(groups):
        if not group:
            continue
        inner = ''.join(eff(sh, 'clickEffect' if i == 0 else 'withEffect')
                        for i, sh in enumerate(group))
        first_cond = ('<p:cond delay="indefinite"/>'
                      '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
                      if gi == 0 else '<p:cond delay="indefinite"/>')
        par_list.append(
            '<p:par><p:cTn id="%d" fill="hold"><p:stCondLst>%s</p:stCondLst><p:childTnLst>'
            '<p:par><p:cTn id="%d" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            '<p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
            '</p:childTnLst></p:cTn></p:par>' % (nid(), first_cond, nid(), inner))

    return (
        '<p:timing xmlns:p="%s"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
        '<p:childTnLst>%s</p:childTnLst></p:cTn>'
        '<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    ) % (NS_P, ''.join(par_list))


def inject_animations():
    from lxml import etree
    total = 0
    for idx, slide in enumerate(prs.slides):
        groups = GROUPS.get(idx) or []
        if not groups:
            continue
        el = etree.fromstring(timing_xml(groups))
        slide._element.append(el)
        total += sum(len(g) for g in groups)
    return total


count = inject_animations()
prs.save(OUT)
print('已生成：' + OUT)
print('页数：%d　带动画的形状：%d 个　文件大小 %.2f MB'
      % (len(prs.slides.__iter__.__self__._sldIdLst), count,
         os.path.getsize(OUT) / 1024 / 1024))
